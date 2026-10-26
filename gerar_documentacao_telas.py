from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

BASE_DIR = Path(__file__).resolve().parent
IMAGENS_DIR = BASE_DIR / "prints_telas"


def adicionar_tela(doc, titulo, descricao, elementos, imagem_nome):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(titulo)
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph(descricao)

    for item in elementos:
        doc.add_paragraph(f"• {item}", style="List Bullet")

    imagem = IMAGENS_DIR / imagem_nome
    if imagem.exists():
        doc.add_paragraph("Tela:")
        doc.add_picture(str(imagem), width=Inches(7.5))
    else:
        doc.add_paragraph("Imagem da tela não encontrada.")

    doc.add_paragraph()


def main():
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Manual do Core - AutoPeças")
    title_run.bold = True
    title_run.font.size = Pt(24)

    doc.add_paragraph(
        "Manual de referência do Core desktop Tkinter e de suas telas e funcionalidades. "
        "A aplicação permite manter o catálogo de peças, consultar o estoque, registrar vendas "
        "e acompanhar valores financeiros. O acesso ao sistema exige uma conta de usuário."
    )

    adicionar_tela(
        doc,
        "1. Configuração do primeiro administrador",
        "Na primeira execução sem usuários cadastrados, o sistema solicita a criação da conta administradora inicial.",
        [
            "Informe um nome de usuário com pelo menos 3 caracteres.",
            "A senha deve possuir pelo menos 8 caracteres e ser confirmada no segundo campo.",
            "A conta inicial recebe o perfil Administrador e é criada apenas uma vez.",
            "A senha é armazenada como hash PBKDF2 com salt individual; não é guardada em texto puro.",
        ],
        "primeiro_administrador.png",
    )

    adicionar_tela(
        doc,
        "2. Login",
        "Tela de autenticação apresentada nas inicializações seguintes e após encerrar uma sessão.",
        [
            "Informe usuário e senha cadastrados para entrar.",
            "Credenciais incorretas não abrem o menu do sistema.",
            "A opção Sair da sessão retorna para esta tela sem encerrar o aplicativo.",
        ],
        "login.png",
    )

    adicionar_tela(
        doc,
        "3. Menu principal",
        "Ponto de entrada e navegação para as áreas operacionais e financeiras do sistema.",
        [
            "Cadastro: abre o formulário para incluir ou editar peças.",
            "Consulta: abre a listagem, busca, filtro por categoria e manutenção do catálogo.",
            "Vendas: abre o registro de saída de peças e o histórico recente.",
            "Consulta Financeira: abre os relatórios de vendas e de valor do estoque.",
            "A sessão identifica o usuário conectado; o atalho de gerenciamento aparece somente para administradores.",
            "Sair da sessão retorna à tela de login.",
            "Administradores acessam cadastro, consulta, vendas, relatórios financeiros e gestão de usuários.",
            "Usuários operacionais acessam consulta de estoque em modo leitura e registro de vendas.",
        ],
        "menu_principal.png",
    )

    adicionar_tela(
        doc,
        "4. Cadastro de peças",
        "A tela abre em modo de novo produto, com os campos vazios. Também permite escolher uma peça existente para editar.",
        [
            "Campos: nome, marca, número de fábrica, categoria, quantidade em estoque, preço e status.",
            "Ao entrar pelo menu, o formulário é limpo e fica pronto para um novo cadastro.",
            "Use o seletor Editar produto existente e Carregar para edição para preencher o formulário com uma peça cadastrada.",
            "Use Novo produto para limpar o formulário e abandonar a edição atual.",
            "Salvar valida os dados e inclui uma nova peça ou atualiza a peça selecionada para edição.",
            "O status é calculado automaticamente pela quantidade: zero unidades significa Esgotado; acima de zero significa Disponível.",
            "Limpar apaga os valores digitados no formulário.",
            "A validação exige os campos descritivos e valores numéricos não negativos para quantidade e preço.",
            "A mensagem de status informa o resultado do cadastro ou da atualização.",
        ],
        "cadastro_peca.png",
    )

    adicionar_tela(
        doc,
        "5. Consulta de peças",
        "Tela de pesquisa, acompanhamento e manutenção do catálogo de peças.",
        [
            "Busca por nome, marca, número de fábrica ou categoria.",
            "Filtro de categoria preenchido a partir das categorias cadastradas; é possível combinar busca e categoria.",
            "A tabela apresenta nome, marca, número de fábrica, categoria, quantidade, preço e status.",
            "O resumo mostra quantidade de itens, unidades em estoque, peças críticas e valor estimado do estoque.",
            "Editar carrega a peça selecionada no formulário de cadastro.",
            "Excluir solicita confirmação antes de remover a peça.",
            "As cores da tabela distinguem apenas itens disponíveis e esgotados, conforme a quantidade.",
        ],
        "consulta_pecas.png",
    )

    adicionar_tela(
        doc,
        "6. Registro de vendas",
        "Registra saídas de peças, atualiza o estoque e mantém uma lista das vendas recentes.",
        [
            "A seleção de produto exige uma consulta; não há uma lista suspensa pré-carregada.",
            "Clique em Consultar produto, pesquise por nome, marca, número de fábrica ou categoria e execute a consulta.",
            "Escolha um resultado com unidades disponíveis e confirme em Selecionar produto.",
            "Informe uma quantidade inteira maior que zero para registrar a venda.",
            "O sistema impede a venda de quantidade superior ao estoque disponível.",
            "Cada venda registra o identificador do produto, nome de referência, quantidade, preço unitário praticado, valor total, data e usuário responsável.",
            "A data é gravada no formato AAAA-MM-DD.",
            "O histórico identifica o usuário que registrou a venda e apresenta o preço unitário.",
            "Vendas históricas sem vínculos legados permanecem preservadas e aparecem como histórico legado.",
            "Após a venda, o status fica Esgotado se o saldo chegar a zero; com qualquer saldo positivo, fica Disponível.",
            "O resumo mostra o valor atual do estoque e o total acumulado de vendas.",
            "O histórico exibe as dez vendas mais recentes.",
        ],
        "registro_vendas.png",
    )

    adicionar_tela(
        doc,
        "7. Consulta financeira",
        "Apresenta vendas agregadas por produto e valor atual do estoque agrupado por categoria.",
        [
            "A busca por produto pode ser combinada com as datas inicial e final.",
            "As datas de filtro devem usar o formato AAAA-MM-DD; sem datas, o relatório considera todo o histórico.",
            "A aba Vendas por produto apresenta unidades vendidas e valor vendido no filtro aplicado.",
            "A aba Estoque por categoria apresenta unidades e valor estimado por categoria.",
            "O resumo exibe o valor vendido no filtro e o valor atual total do estoque.",
            "Vendas antigas sem data aparecem no relatório geral, mas não podem ser atribuídas a um período.",
        ],
        "consulta_financeira.png",
    )

    adicionar_tela(
        doc,
        "8. Gerenciamento de usuários",
        "Módulo administrativo para criar e remover contas de acesso ao sistema.",
        [
            "A tela é acessível somente a usuários com perfil Administrador.",
            "O administrador cadastra usuários operacionais com senha de pelo menos 8 caracteres.",
            "A tabela lista usuário, perfil e data de criação.",
            "A remoção exige confirmação e não permite apagar a própria conta conectada.",
            "A remoção do último administrador é bloqueada para evitar perder o controle das contas.",
            "Todos os usuários compartilham o mesmo estoque e histórico de vendas local.",
        ],
        "gerenciamento_usuarios.png",
    )

    doc.add_heading("9. Fluxos de uso", level=1)
    doc.add_paragraph("Acesso ao sistema")
    fluxo_acesso = doc.add_table(rows=1, cols=2)
    fluxo_acesso.style = "Table Grid"
    fluxo_acesso.rows[0].cells[0].text = "Passo"
    fluxo_acesso.rows[0].cells[1].text = "Ação"
    for passo, acao in [
        ("1", "Na primeira execução, crie a conta administradora inicial."),
        ("2", "Nas próximas execuções, entre com usuário e senha."),
        ("3", "O administrador pode cadastrar ou remover contas no módulo de usuários."),
        ("4", "Encerre a sessão pelo botão Sair da sessão."),
    ]:
        row = fluxo_acesso.add_row().cells
        row[0].text = passo
        row[1].text = acao

    doc.add_paragraph("Cadastro e manutenção de peças")
    fluxo = doc.add_table(rows=1, cols=2)
    fluxo.style = "Table Grid"
    header_cells = fluxo.rows[0].cells
    header_cells[0].text = "Passo"
    header_cells[1].text = "Ação"

    dados_fluxo = [
        ("1", "Acesse Cadastro pelo menu principal."),
        ("2", "Preencha os dados obrigatórios e salve a peça."),
        ("3", "Use Consulta para buscar, filtrar, editar ou excluir itens."),
        ("4", "Confirme a exclusão quando solicitado."),
    ]

    for passo, acao in dados_fluxo:
        row = fluxo.add_row().cells
        row[0].text = passo
        row[1].text = acao

    doc.add_paragraph("Registro e análise de vendas")
    fluxo_vendas = doc.add_table(rows=1, cols=2)
    fluxo_vendas.style = "Table Grid"
    fluxo_vendas.rows[0].cells[0].text = "Passo"
    fluxo_vendas.rows[0].cells[1].text = "Ação"
    for passo, acao in [
        ("1", "Acesse Vendas, selecione a peça e informe a quantidade."),
        ("2", "Registre a venda; o estoque é reduzido e o histórico atualizado."),
        ("3", "Acesse Consulta Financeira para filtrar as vendas por produto ou período."),
        ("4", "Consulte o valor atual do estoque agrupado por categoria."),
    ]:
        row = fluxo_vendas.add_row().cells
        row[0].text = passo
        row[1].text = acao

    doc.add_heading("Persistência de dados", level=1)
    doc.add_paragraph(
        "O sistema utiliza o banco SQLite local autopecas.db para armazenar peças e vendas. "
        "Na primeira execução, os arquivos JSON legados são importados uma única vez, quando disponíveis. "
        "Os arquivos JSON originais são preservados como backup. O registro de uma venda e a atualização do estoque "
        "são confirmados juntos em uma transação."
    )

    doc.add_heading("Histórico de alterações", level=1)
    historico = doc.add_table(rows=1, cols=2)
    historico.style = "Table Grid"
    historico.rows[0].cells[0].text = "Data"
    historico.rows[0].cells[1].text = "Alteração"
    row = historico.add_row().cells
    row[0].text = "2026-10-23"
    row[1].text = (
        "Cadastro de produtos passa a abrir com o formulário vazio. Foi adicionado seletor para carregar "
        "um produto existente para edição e botão Novo produto para iniciar ou retornar ao modo de inclusão."
    )
    row = historico.add_row().cells
    row[0].text = "2026-10-23"
    row[1].text = (
        "Status do produto passa a ser calculado pela quantidade em estoque: zero é Esgotado e quantidade positiva "
        "é Disponível. A opção Em breve foi removida e status legado no SQLite é normalizado ao iniciar o sistema."
    )
    row = historico.add_row().cells
    row[0].text = "2026-10-23"
    row[1].text = (
        "Na tela de vendas, a seleção por lista suspensa foi substituída por consulta explícita de produtos "
        "por nome, marca, número de fábrica ou categoria, com seleção em tabela de resultados."
    )
    row = historico.add_row().cells
    row[0].text = "2026-10-23"
    row[1].text = (
        "Acesso operacional separado do administrativo. Vendas novas passam a registrar produto relacionado, "
        "preço unitário praticado e usuário responsável; histórico legado é preservado sem atribuições inventadas."
    )
    doc.save(BASE_DIR / "documentacao_core.docx")
    print("Arquivo gerado: documentacao_core.docx")


if __name__ == "__main__":
    main()
